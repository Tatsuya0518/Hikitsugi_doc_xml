import io
import os
import base64
import xml.etree.ElementTree as ET

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Color


def _parse_color(color_str: str | None) -> Color | None:
    """Convert an 8-char AARRGGBB (or 6-char RRGGBB) hex string to openpyxl Color."""
    if not color_str:
        return None
    s = color_str.lstrip("#")
    if len(s) == 6:
        s = "FF" + s
    if len(s) == 8:
        return Color(rgb=s)
    return None


def _border_side(style_str: str | None) -> Side:
    """Map a border-style name to an openpyxl Side."""
    if not style_str or style_str == "none":
        return Side(border_style=None)
    return Side(border_style=style_str)


def export_xml_to_excel(xml_path: str, output_path: str) -> str:
    """
    Convert the custom XML format back to an Excel (.xlsx) file,
    restoring values, formatting, merged cells, and embedded images.
    """
    tree = ET.parse(xml_path)
    root = tree.getroot()

    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # remove the default blank sheet

    for sheet_elem in root.findall("sheet"):
        sheet_name = sheet_elem.get("name", "Sheet")
        ws = wb.create_sheet(title=sheet_name)

        # ── Column widths ──────────────────────────────────────────────────────
        for col_el in sheet_elem.findall("colWidths/col"):
            idx = int(col_el.get("index", 1))
            width = float(col_el.get("width", 8.43))
            letter = get_column_letter(idx)
            ws.column_dimensions[letter].width = width
            if col_el.get("hidden") == "true":
                ws.column_dimensions[letter].hidden = True

        # ── Row heights ────────────────────────────────────────────────────────
        for row_el in sheet_elem.findall("rowHeights/row"):
            idx = int(row_el.get("index", 1))
            height = float(row_el.get("height", 15))
            ws.row_dimensions[idx].height = height
            if row_el.get("hidden") == "true":
                ws.row_dimensions[idx].hidden = True

        # ── Cells ──────────────────────────────────────────────────────────────
        for row_el in sheet_elem.findall("rows/row"):
            for cell_el in row_el.findall("cell"):
                r = int(cell_el.get("row", 1))
                c = int(cell_el.get("col", 1))
                value = cell_el.get("value", "")
                cell_type = cell_el.get("type", "s")

                cell = ws.cell(row=r, column=c)

                if value:
                    if cell_type == "n":
                        try:
                            cell.value = float(value)
                        except ValueError:
                            cell.value = value
                    elif cell_type == "b":
                        cell.value = value.lower() == "true"
                    else:
                        cell.value = value

                fmt = cell_el.find("format")
                if fmt is not None:
                    # Font
                    font_color = _parse_color(fmt.get("fontColor", ""))
                    cell.font = Font(
                        bold=(fmt.get("bold") == "true"),
                        italic=(fmt.get("italic") == "true"),
                        underline=("single"
                                   if fmt.get("underline") == "true"
                                   else None),
                        size=float(fmt.get("fontSize", 11) or 11),
                        name=(fmt.get("fontName") or "Calibri"),
                        color=font_color,
                    )

                    # Fill
                    bg_str = fmt.get("bgColor", "")
                    if bg_str:
                        bg_color = _parse_color(bg_str)
                        if bg_color:
                            cell.fill = PatternFill(
                                fill_type="solid", fgColor=bg_color
                            )

                    # Alignment
                    cell.alignment = Alignment(
                        horizontal=(fmt.get("halign") or None),
                        vertical=(fmt.get("valign") or None),
                        wrap_text=(fmt.get("wrapText") == "true"),
                        indent=int(fmt.get("indent", 0) or 0),
                    )

                    # Border
                    cell.border = Border(
                        top=_border_side(fmt.get("borderTop")),
                        bottom=_border_side(fmt.get("borderBottom")),
                        left=_border_side(fmt.get("borderLeft")),
                        right=_border_side(fmt.get("borderRight")),
                    )

                    # Number format
                    num_fmt = fmt.get("numFmt", "General")
                    if num_fmt:
                        cell.number_format = num_fmt

        # ── Merged cells ──────────────────────────────────────────────────────
        for merge_el in sheet_elem.findall("merges/merge"):
            ref = merge_el.get("ref")
            if ref:
                ws.merge_cells(ref)

        # ── Embedded images ───────────────────────────────────────────────────
        for img_el in sheet_elem.findall("images/image"):
            img_b64 = img_el.get("data", "")
            if not img_b64:
                continue
            try:
                from openpyxl.drawing.image import Image as XLImage
                img_data = base64.b64decode(img_b64)
                img_io = io.BytesIO(img_data)
                xl_img = XLImage(img_io)

                width = img_el.get("width")
                height = img_el.get("height")
                if width:
                    xl_img.width = float(width)
                if height:
                    xl_img.height = float(height)

                row = int(img_el.get("row", 1))
                col = int(img_el.get("col", 1))
                cell_ref = f"{get_column_letter(col)}{row}"
                ws.add_image(xl_img, cell_ref)
            except Exception:
                pass  # skip images that fail to decode

    wb.save(output_path)
    return output_path
