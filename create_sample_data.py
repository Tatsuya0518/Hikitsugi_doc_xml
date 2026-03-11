"""
Sample Excel data generator – run once to populate data/watch/ with demo files.
Usage:  python3 create_sample_data.py
"""
import os

import openpyxl
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, Color
)
from openpyxl.utils import get_column_letter

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(BASE_DIR, "data", "watch")


def thin_border(sides="TRBL"):
    thin = Side(border_style="thin", color="FF666666")
    return Border(
        top=thin    if "T" in sides else Side(),
        right=thin  if "R" in sides else Side(),
        bottom=thin if "B" in sides else Side(),
        left=thin   if "L" in sides else Side(),
    )


def header_fill():
    return PatternFill(fill_type="solid", fgColor=Color(rgb="FF4472C4"))


def alt_fill():
    return PatternFill(fill_type="solid", fgColor=Color(rgb="FFE8F0FE"))


def make_handover_doc(wb, sheet_name: str, title: str, rows: list[list]):
    """Create a styled handover sheet."""
    ws = wb.create_sheet(title=sheet_name)

    # Title row (merged A1:D1)
    ws.merge_cells("A1:D1")
    title_cell = ws["A1"]
    title_cell.value = title
    title_cell.font = Font(name="Meiryo", bold=True, size=14, color="FFFFFFFF")
    title_cell.fill = PatternFill(fill_type="solid", fgColor=Color(rgb="FF2F5496"))
    title_cell.alignment = Alignment(horizontal="center", vertical="center",
                                     wrap_text=True)
    ws.row_dimensions[1].height = 30

    # Header row
    headers = ["項目", "内容", "担当者", "更新日"]
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=2, column=col, value=h)
        cell.font = Font(name="Meiryo", bold=True, size=11, color="FFFFFFFF")
        cell.fill = header_fill()
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border()
    ws.row_dimensions[2].height = 22

    # Data rows
    for row_idx, row_data in enumerate(rows, start=3):
        for col_idx, value in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = Font(name="Meiryo", size=10)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = thin_border()
            if row_idx % 2 == 0:
                cell.fill = alt_fill()
        ws.row_dimensions[row_idx].height = 18

    # Column widths
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 40
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 14

    return ws


def create_team1_aaa01():
    path = os.path.join(DATA_DIR, "チーム1", "AAA-01")
    os.makedirs(path, exist_ok=True)
    outfile = os.path.join(path, "AAA-01作業引継ぎドキュメント.xlsx")

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    rows = [
        ["作業概要",   "本番環境のデプロイ手順および注意事項をまとめたドキュメントです。",  "田中", "2024-04-01"],
        ["前提条件",   "SSH鍵の設定・VPN接続・必要権限の取得が完了していること。",         "田中", "2024-04-01"],
        ["デプロイ手順", "1. gitリポジトリからmainブランチをpull\n2. docker-compose up -d",    "山田", "2024-05-10"],
        ["ロールバック", "問題発生時は前バージョンのタグをcheckoutしてdeployを実行。",        "山田", "2024-05-10"],
        ["連絡先",     "障害時は#ops-alertチャンネルへ。深夜は携帯090-XXXX-XXXXへ。",       "佐藤", "2024-06-01"],
    ]
    make_handover_doc(wb, "デプロイ手順", "AAA-01 デプロイ作業引継ぎ", rows)

    # Second sheet: checklist
    ws2 = wb.create_sheet(title="確認チェックリスト")
    ws2.merge_cells("A1:C1")
    ws2["A1"].value = "デプロイ前確認チェックリスト"
    ws2["A1"].font = Font(bold=True, size=12, color="FF000000")
    ws2["A1"].fill = PatternFill(fill_type="solid",
                                  fgColor=Color(rgb="FFFFFFAA"))
    ws2["A1"].alignment = Alignment(horizontal="center")
    ws2.row_dimensions[1].height = 24
    checks = [
        ["確認項目",       "確認結果", "備考"],
        ["全テスト通過",   "□",        ""],
        ["ステージング確認", "□",       ""],
        ["承認取得",       "□",        ""],
        ["バックアップ取得", "□",       ""],
        ["通知送信済み",   "□",        ""],
    ]
    for r, row in enumerate(checks, start=2):
        for c, val in enumerate(row, start=1):
            cell = ws2.cell(row=r, column=c, value=val)
            cell.border = thin_border()
            if r == 2:
                cell.font = Font(bold=True)
                cell.fill = header_fill()
                cell.font = Font(bold=True, color="FFFFFFFF")
            else:
                cell.font = Font(name="Meiryo", size=10)
    ws2.column_dimensions["A"].width = 24
    ws2.column_dimensions["B"].width = 12
    ws2.column_dimensions["C"].width = 30

    wb.save(outfile)
    print(f"Created: {outfile}")


def create_team2_bbb01():
    path = os.path.join(DATA_DIR, "チーム2", "BBB-01")
    os.makedirs(path, exist_ok=True)
    outfile = os.path.join(path, "BBB-01作業引継ぎドキュメント.xlsx")

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    rows = [
        ["DB接続情報",  "host: db.example.com / port: 5432 / db: app_prod",    "鈴木", "2024-03-15"],
        ["バックアップ", "毎日02:00 JST に S3 へ自動バックアップ。保持期間30日。", "鈴木", "2024-03-15"],
        ["マイグレーション", "alembic upgrade head を実行後、再起動すること。",    "伊藤", "2024-04-20"],
        ["定期メンテ",  "毎月第1土曜の深夜2時～4時がメンテナンスウィンドウ。",    "伊藤", "2024-04-20"],
        ["監視設定",   "Datadog ダッシュボード URL: https://app.datadoghq.com/…", "高橋", "2024-05-01"],
    ]
    make_handover_doc(wb, "DB管理手順", "BBB-01 データベース管理引継ぎ", rows)
    wb.save(outfile)
    print(f"Created: {outfile}")


def create_team2_cbaa01():
    path = os.path.join(DATA_DIR, "チーム2", "CBA-A-01")
    os.makedirs(path, exist_ok=True)
    outfile = os.path.join(path, "CBA-A-01作業引継ぎドキュメント.xlsx")

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    rows = [
        ["サービス概要",  "社内認証基盤。SSOを提供しSSL証明書は Let's Encrypt で自動更新。", "渡辺", "2024-02-01"],
        ["証明書更新",   "certbot renew を月1実行。失敗時はメールアラート。",              "渡辺", "2024-02-01"],
        ["ユーザ追加",   "ldapadd コマンドで追加後、管理画面から同期ボタンを押す。",        "中村", "2024-03-10"],
        ["ログ確認",     "/var/log/auth.log を確認。異常ログはインシデント管理へ。",        "中村", "2024-03-10"],
        ["緊急連絡",     "セキュリティインシデントは security@example.com へ即報。",       "小林", "2024-05-20"],
    ]
    make_handover_doc(wb, "認証基盤手順", "CBA-A-01 認証基盤引継ぎ", rows)
    wb.save(outfile)
    print(f"Created: {outfile}")


if __name__ == "__main__":
    create_team1_aaa01()
    create_team2_bbb01()
    create_team2_cbaa01()
    print("サンプルデータ作成完了。")
